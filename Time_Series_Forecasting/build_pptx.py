# -*- coding: utf-8 -*-
"""
Builds Time_Series_Forecasting_Beginner_Course.pptx

Run from the project root:
    python build_pptx.py
"""

from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx_helpers import *

V = "visuals/"

prs = new_presentation()

CONTENT_TOP = Inches(1.55)
LEFT_X = Inches(0.78)
FULL_W = Inches(11.78)
COL1_X = Inches(0.78)
COL1_W = Inches(5.55)
COL2_X = Inches(6.60)
COL2_W = Inches(6.00)


def slide_bullets(kicker, title, section, bullets, notes="", subtitle=None, size=17.5, marker="•  "):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    top = CONTENT_TOP
    if subtitle:
        add_text(s, LEFT_X, top, FULL_W, Inches(0.5), subtitle, size=16, color=MUTED, italic=True)
        top = Emu(top + Inches(0.55))
    add_bullets(s, LEFT_X, top, FULL_W, Inches(6.95) - top, bullets, size=size, marker=marker)
    if notes:
        add_notes(s, notes)
    return s


def slide_image_right(kicker, title, section, bullets, image, notes="", caption=None,
                       img_width=Inches(5.85), subtitle=None):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    top = CONTENT_TOP
    if subtitle:
        add_text(s, COL1_X, top, COL1_W, Inches(0.5), subtitle, size=15.5, color=MUTED, italic=True)
        top = Emu(top + Inches(0.55))
    add_bullets(s, COL1_X, top, COL1_W, Inches(6.9) - top, bullets, size=16.5)
    add_picture_framed(s, V + image, COL2_X, Inches(1.85), w=img_width, caption=caption)
    if notes:
        add_notes(s, notes)
    return s


def slide_image_full(kicker, title, section, image, notes="", caption=None,
                      img_width=Inches(9.4), intro=None, y=Inches(1.75)):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    if intro:
        add_text(s, LEFT_X, CONTENT_TOP, FULL_W, Inches(0.5), intro, size=16.5, color=MUTED, italic=True)
        y = Emu(y + Inches(0.35))
    x = Emu(int((SLIDE_W - img_width) / 2))
    add_picture_framed(s, V + image, x, y, w=img_width, caption=caption)
    if notes:
        add_notes(s, notes)
    return s


def slide_flow_v(kicker, title, section, steps, notes="", intro=None, colors=None, box_w=Inches(5.6)):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    y = CONTENT_TOP
    if intro:
        add_text(s, LEFT_X, y, FULL_W, Inches(0.5), intro, size=16.5, color=MUTED, italic=True)
        y = Emu(y + Inches(0.5))
    x = Emu(int((SLIDE_W - box_w) / 2))
    n = len(steps)
    if n > 8:
        box_h, gap, fsize = Inches(0.44), Inches(0.09), 13
    elif n > 6:
        box_h, gap, fsize = Inches(0.5), Inches(0.16), 14.5
    elif n > 5:
        box_h, gap, fsize = Inches(0.56), Inches(0.20), 15
    else:
        box_h, gap, fsize = Inches(0.62), Inches(0.26), 15
    if colors is None:
        cy = y
        for i, step in enumerate(steps):
            add_box_with_text(s, x, cy, box_w, box_h, step, fill=BLUE, radius=0.18, size=fsize)
            cy = Emu(cy + box_h + gap)
            if i < n - 1:
                cx = Emu(x + box_w // 2)
                add_arrow_between(s, cx, Emu(cy - gap + Pt(2)), cx, Emu(cy - Pt(2)))
    else:
        cy = y
        for i, (step, c) in enumerate(zip(steps, colors)):
            add_box_with_text(s, x, cy, box_w, box_h, step, fill=c, radius=0.18, size=fsize)
            cy = Emu(cy + box_h + gap)
            if i < n - 1:
                cx = Emu(x + box_w // 2)
                add_arrow_between(s, cx, Emu(cy - gap + Pt(2)), cx, Emu(cy - Pt(2)))
    if notes:
        add_notes(s, notes)
    return s


def slide_flow_h(kicker, title, section, steps, notes="", intro=None, colors=None, y=Inches(3.1), h=Inches(1.0)):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    if intro:
        add_text(s, LEFT_X, CONTENT_TOP, FULL_W, Inches(0.5), intro, size=16.5, color=MUTED, italic=True)
    add_flow_horizontal(s, steps, LEFT_X, y, FULL_W, h=h, colors=colors)
    if notes:
        add_notes(s, notes)
    return s


def slide_table(kicker, title, section, headers, rows, notes="", intro=None,
                 col_widths=None, table_h=Inches(3.6), header_fill=INK):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    top = CONTENT_TOP
    if intro:
        add_text(s, LEFT_X, top, FULL_W, Inches(0.5), intro, size=16, color=MUTED, italic=True)
        top = Emu(top + Inches(0.55))
    add_table(s, LEFT_X, top, FULL_W, table_h, headers, rows, col_widths=col_widths, header_fill=header_fill)
    if notes:
        add_notes(s, notes)
    return s


def slide_code(kicker, title, section, code, bullets=None, notes="", code_h=Inches(2.2), intro=None):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    top = CONTENT_TOP
    if intro:
        add_text(s, LEFT_X, top, FULL_W, Inches(0.45), intro, size=16, color=MUTED, italic=True)
        top = Emu(top + Inches(0.5))
    add_code_block(s, LEFT_X, top, FULL_W, code_h, code)
    if bullets:
        bullet_top = Emu(top + code_h + Inches(0.3))
        bullet_h = Emu(int(Inches(6.9)) - int(bullet_top))
        add_bullets(s, LEFT_X, bullet_top, FULL_W, bullet_h, bullets, size=15.5)
    if notes:
        add_notes(s, notes)
    return s


def slide_quote(kicker, title, section, quote, subtext=None, notes="", quote_size=32, fill=DARK_PANEL,
                 text_color=WHITE):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    panel = add_rect(s, LEFT_X, Inches(2.1), FULL_W, Inches(3.6), fill=fill, radius=0.06)
    tf = panel.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(30); tf.margin_right = Pt(30)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = quote
    set_font(run, quote_size, text_color, True)
    if subtext:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(14)
        run2 = p2.add_run()
        run2.text = subtext
        set_font(run2, 17, RGBColor(0xC7, 0xCE, 0xD9), False, True)
    if notes:
        add_notes(s, notes)
    return s


SEC1 = "Section 1 · Introduction"
SEC2 = "Section 2 · The Real Problem"
SEC3 = "Section 3 · What Is Time Series Data?"
SEC4 = "Section 4 · Understanding Patterns"
SEC5 = "Section 5 · Visualize First"
SEC6 = "Section 6 · Lags"
SEC7 = "Section 7 · Moving Averages"
SEC8 = "Section 8 · Forecasting"
SEC9 = "Section 9 · Train / Test Split"
SEC10 = "Section 10 · Evaluation"
SEC11 = "Section 11 · Simple Models"
SEC12 = "Section 12 · Business Interpretation"

# ===========================================================================
# SLIDE 1 — TITLE
# ===========================================================================
s = blank_slide(prs, bg=DARK_PANEL)
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(0.14), fill=TEAL)
add_text(s, Inches(0.9), Inches(1.15), Inches(6), Inches(0.4), "BEGINNER DATA SCIENCE COURSE",
          size=14, color=RGBColor(0x99, 0xF6, 0xE4), bold=True)
add_text(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.3), "Time Series Forecasting", size=46, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(3.25), Inches(11.5), Inches(1.1), "Predicting the Future Using Historical Data",
          size=28, color=RGBColor(0x99, 0xF6, 0xE4), bold=True)
add_rect(s, Inches(0.9), Inches(4.5), Inches(1.3), Pt(3.5), fill=AMBER)
add_text(s, Inches(0.9), Inches(4.7), Inches(10.5), Inches(0.6),
          "Understanding Time Series Data, Patterns, and Forecasting", size=19, color=RGBColor(0xE5, 0xE7, 0xEB), italic=True)
for i, (label, c) in enumerate([("Past Data", MUTED), ("Find Patterns", BLUE), ("Forecast Forward", AMBER)]):
    add_pill(s, Inches(0.9 + i * 2.55), Inches(6.1), Inches(2.35), Inches(0.5), label, fill=c, text_color=WHITE, size=13)
add_notes(s, "Welcome the class. This is a 2-hour session: Hour 1 is concepts (this deck), Hour 2 is hands-on "
             "Python. Frame it as building ONE forecasting system together for a bank's daily cash withdrawals, "
             "not a tour of separate techniques. Ask: has anyone ever tried to guess tomorrow's weather, traffic, "
             "or a store's busy period based on what usually happens? That everyday instinct — using the past to "
             "guess the future — is exactly what we're formalizing today.")

# ===========================================================================
# SLIDE 2 — LEARNING OBJECTIVES
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, "Introduction", "What Will We Learn?", section="Slide 2")
left_items = ["What Time Series data is", "Why time order matters", "Trend", "Seasonality", "Cycles", "Noise",
              "Lag", "Moving averages"]
right_items = ["Forecasting", "Time-based train/test split", "Baseline forecasting", "A simple forecasting model",
               "Forecast evaluation (MAE, RMSE, MAPE)", "Building a forecast in Python", "Data leakage",
               "Turning a forecast into a business decision"]
add_bullets(s, COL1_X, CONTENT_TOP, COL1_W, Inches(5.3), left_items, size=16.5, space_after=11)
add_bullets(s, COL2_X, CONTENT_TOP, COL2_W, Inches(5.3), right_items, size=15.5, space_after=11)
add_notes(s, "Walk down the list quickly — this is a map, not a lesson. Mention explicitly that we are "
             "deliberately NOT covering ARIMA mathematics, SARIMA, Prophet, or deep learning today — those are "
             "later-course territory. Today is entirely about building solid intuition and one working, simple "
             "forecasting pipeline.")

# ===========================================================================
# SLIDE 3 — START WITH A REAL BUSINESS QUESTION
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC2, "Start With a Real Business Question", section="Slide 3")
add_table(s, LEFT_X, Inches(1.75), Inches(7.2), Inches(2.4),
          ["Date", "Withdrawal Amount"],
          [["Monday", "₦25,000,000"], ["Tuesday", "₦28,000,000"], ["Wednesday", "₦30,000,000"],
           ["Thursday", "₦27,000,000"], ["Friday", "₦45,000,000"]],
          col_widths=[1, 1], header_fill=INK)
add_text(s, Inches(8.3), Inches(2.0), Inches(3.6), Inches(3.0),
          "A bank records daily ATM withdrawals.\n\nCan we predict tomorrow's withdrawal demand?\n\n"
          "Can we predict next week's demand?", size=16.5, color=SUBTEXT)
add_notes(s, "This is the running example for the ENTIRE course — a bank forecasting daily cash demand. Let the "
             "class react to the jump on Friday before explaining anything. Ask: 'if you were the bank's cash "
             "operations manager, why would knowing tomorrow's number matter?' (Answer, if needed: too little "
             "cash = angry customers and reputational damage; too much idle cash = wasted opportunity cost and "
             "security risk.)")

# ===========================================================================
# SLIDE 4 — WHAT MAKES THIS DIFFERENT?
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC2, "What Makes This Different?", section="Slide 4")
add_text(s, COL1_X, Inches(1.85), COL1_W, Inches(0.5), "Normal dataset:", size=15.5, color=MUTED, bold=True)
add_box_with_text(s, COL1_X, Inches(2.35), COL1_W, Inches(2.1), "Row 1\nRow 2\nRow 3\n\n(order may not matter)",
                   fill=BG_SOFT, text_color=INK, size=16, bold=False, line_color=LINE)
add_text(s, COL2_X, Inches(1.85), COL2_W, Inches(0.5), "Time Series dataset:", size=15.5, color=MUTED, bold=True)
add_flow_vertical(s, ["Yesterday", "Today", "Tomorrow"], COL2_X, Inches(2.35), COL2_W, box_h=Inches(0.6), gap=Inches(0.22), fill=TEAL)
add_pill(s, Inches(2.9), Inches(5.85), Inches(7.5), Inches(0.65),
          "Time is one of the most important pieces of information here", fill=BLUE_SOFT, text_color=BLUE_DARK, size=15)
add_notes(s, "This contrast is worth slowing down for — it's the single strongest concept in the whole course. "
             "In a normal supervised ML dataset, shuffling rows changes nothing about what the model can learn. "
             "In a time series, shuffling destroys the very information (the sequence) we're trying to use. Say "
             "the word TIME MATTERS out loud, more than once.")

# ===========================================================================
# SLIDE 5 — DEFINITION
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC3, "What Is Time Series Data?", "Slide 5")
add_text(s, LEFT_X, Inches(1.85), FULL_W, Inches(0.9),
          "“A Time Series is a sequence of observations\nrecorded over time.”", size=22, color=INK, bold=True)
examples = ["Daily sales", "Monthly revenue", "Stock prices", "Bank transactions", "Temperature",
            "Website visitors", "Exchange rates", "Electricity demand", "Cash withdrawals"]
for i, ex in enumerate(examples):
    x = LEFT_X + Inches((i % 3) * 3.95)
    y = Inches(3.3) + Inches((i // 3) * 0.85)
    add_pill(s, x, y, Inches(3.7), Inches(0.6), ex, fill=BLUE_SOFT, text_color=BLUE_DARK, size=14)
add_notes(s, "Ask learners to shout out their own example before you reveal these. Almost anything measured "
             "repeatedly over time qualifies — that's the point. The bank's ATM withdrawal series is just one "
             "instance of a pattern that shows up across virtually every industry, which the next slide makes "
             "explicit.")

# ===========================================================================
# SLIDE 6 — REAL-LIFE EXAMPLES
# ===========================================================================
slide_table(SEC3, "Real-Life Examples", "Slide 6",
            ["Industry", "Forecasting Example"],
            [["Banking", "Cash demand"], ["Retail", "Product sales"], ["Healthcare", "Patient visits"],
             ["Transport", "Passenger demand"], ["Energy", "Electricity consumption"],
             ["Agriculture", "Crop production"], ["Technology", "Website traffic"]],
            col_widths=[1, 1.4], table_h=Inches(3.9),
            notes="Emphasize: the same forecasting PRINCIPLES apply across every one of these rows — trend, "
                  "seasonality, lag, train/test split, evaluation. Once you've learned to forecast bank cash "
                  "demand, you have 80% of what you need to forecast hospital patient visits or electricity "
                  "consumption too.")

# ===========================================================================
# SLIDE 7 — A TIME SERIES IS NOT JUST NUMBERS
# ===========================================================================
slide_image_full(SEC4, "A Time Series Is Not Just Numbers", "Slide 7", "06_decomposition_panel.png",
                  intro="Real-world data often contains several patterns layered on top of each other.",
                  img_width=Inches(6.3), y=Inches(1.95),
                  notes="This chart splits our own ATM series into its three layers: a smooth Trend, a repeating "
                        "weekly Seasonality (zoomed in so the wiggle is visible), and leftover Noise. Don't "
                        "explain the math behind this decomposition — just use it to preview the next four "
                        "slides, which cover Trend, Seasonality, Cycles, and Noise one at a time.")

# ===========================================================================
# SLIDE 8 — TREND
# ===========================================================================
slide_image_right(SEC4, "Trend", "Slide 8",
                   ["“A trend shows the general direction of data over a long period of time.”",
                    "Can be upward, downward, or show no clear direction.",
                    "Our bank's withdrawals trend gently upward over the two years — more customers, more "
                    "cash activity over time.",
                    "A trend tells you WHERE the series is heading in general, not what happens day to day."],
                   "02_trend.png",
                   notes="Point at the red 30-day trend line riding on top of the noisy daily data — that's "
                         "exactly what a trend represents: the slow-moving signal underneath the day-to-day "
                         "noise. Ask: 'is this trend obvious from the raw daily line alone, or does it help to "
                         "smooth it first?' (Preview of moving averages, coming in Section 7.)")

# ===========================================================================
# SLIDE 9 — SEASONALITY
# ===========================================================================
slide_image_right(SEC4, "Seasonality", "Slide 9",
                   ["“Seasonality is a pattern that repeats at regular intervals.”",
                    "Higher sales during Christmas.", "Higher ATM withdrawals at month-end (payday).",
                    "Higher transactions on Fridays.", "Higher electricity use during hot seasons.",
                    "In our data: withdrawals spike on Fridays and drop sharply every Sunday, every single week."],
                   "03_seasonality_weekly.png",
                   notes="This chart is the average withdrawal BY day of week, across the entire 2-year dataset "
                         "— the pattern is remarkably consistent. Follow up with the zoomed-in 6-week chart "
                         "(next few slides in the notebook) if you want to show the SAME shape repeating "
                         "visually over real dates, not just an average.")

# ===========================================================================
# SLIDE 10 — CYCLES
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC4, "Cycles", "Slide 10")
add_text(s, LEFT_X, Inches(1.85), FULL_W, Inches(1.0),
          "Cycles are patterns that rise and fall over longer periods, but may NOT repeat at\nperfectly regular intervals.",
          size=18.5, color=INK, bold=True)
add_box_with_text(s, LEFT_X, Inches(3.1), Inches(5.6), Inches(2.2), "Seasonality\n\nUsually follows a regular,\n"
                   "predictable interval\n(every week, every month)", fill=TEAL, size=16, radius=0.08)
add_box_with_text(s, Inches(6.85), Inches(3.1), Inches(5.6), Inches(2.2), "Cycles\n\nMay not occur at fixed\n"
                   "intervals (e.g. economic\nboom-and-bust cycles)", fill=PURPLE, size=16, radius=0.08)
add_notes(s, "Keep this simple, as instructed — cycles are a beginner-relevant concept but not a slide worth "
             "dwelling on. Example: an economic recession affecting bank withdrawals might last 18 months one "
             "time and 3 years the next — unlike seasonality, you can't set your watch by it. Our own dataset "
             "does not contain a strong cycle (2 years is too short to show one) — flag that honestly if asked.")

# ===========================================================================
# SLIDE 11 — NOISE
# ===========================================================================
slide_image_right(SEC4, "Noise", "Slide 11",
                   ["“Noise represents random variation that is difficult to predict.”",
                    "An unexpected ATM failure.", "An unexpected public holiday.",
                    "An unexpected economic event.",
                    "Not every movement in a Time Series can be predicted — and that's OK."],
                   "05_noise_holiday.png",
                   notes="This chart shows a real public holiday (Oct 1) in our dataset — withdrawals spike the "
                         "day BEFORE (people stocking up on cash) and drop sharply ON the holiday itself (most "
                         "branches closed). A model that doesn't know about holidays will treat this as pure "
                         "noise; one important lesson for later is that some 'noise' can actually be explained "
                         "with the right feature (a holiday indicator, which our dataset includes).")

# ===========================================================================
# SLIDE 12 — FIRST RULE OF TIME SERIES
# ===========================================================================
slide_quote(SEC5, "The First Rule of Time Series", "Slide 12",
            "“Before forecasting, visualize the data.”",
            subtext="Visualization helps you spot trend, seasonality, spikes, outliers, missing periods, and sudden changes.",
            notes="Make this genuinely memorable — some instructors have learners repeat it back. This is not "
                  "just good practice, it's a direct segue into the hands-on notebook, where Step 1 after "
                  "loading data is always 'plot it and look at it' before any modeling begins.")

# ===========================================================================
# SLIDE 13 — LOOK AT THE DATA
# ===========================================================================
slide_image_full(SEC5, "Look at the Data", "Slide 13", "01_full_series.png",
                  intro="What patterns can you see? Look for: increase, decrease, a repeating shape, sudden spikes.",
                  img_width=Inches(10.6), y=Inches(2.0),
                  notes="Let learners call out what they see BEFORE you explain it — most will spot the upward "
                        "drift and the regular spiky rhythm without prompting, and a few will notice the taller "
                        "spikes (holidays) scattered through the series. Do not reveal the trend/seasonality "
                        "explanation yet if you haven't already — let this be a genuine 'what do YOU see' moment.")

# ===========================================================================
# SLIDE 14 — WHAT IS A LAG?
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC6, "What Is a Lag?", "Slide 14")
add_flow_vertical(s, ["Today", "Look Back", "Yesterday", "Last Week", "Last Month"],
                   Inches(4.0), Inches(1.75), Inches(5.3), box_h=Inches(0.6), gap=Inches(0.2), fill=BLUE)
add_text(s, LEFT_X, Inches(5.95), FULL_W, Inches(0.9),
          "“A lag is a previous value of a variable.”\ne.g. Sales Today, Sales Yesterday, Sales 7 Days Ago, Sales 30 Days Ago",
          size=15.5, color=SUBTEXT, italic=True, align=PP_ALIGN.CENTER)
add_notes(s, "Keep this concrete: a 'lag' is just a fancy word for 'a value from some point in the past, lined "
             "up next to today's row.' We'll turn this into actual code (`.shift()`) in the next hour. Preview: "
             "Lag_1 = yesterday, Lag_7 = same day last week.")

# ===========================================================================
# SLIDE 15 — WHY LAGS MATTER
# ===========================================================================
slide_image_right(SEC6, "Why Lags Matter", "Slide 15",
                   ["“Yesterday's sales may contain useful information about today's sales.”",
                    "Today's withdrawal total may depend partly on yesterday's, last week's, and last month's totals.",
                    "This is called TIME DEPENDENCY — values close together in time tend to be related.",
                    "Lag features are how we hand that time dependency to a regular Machine Learning model."],
                   "11_lag_concept.png",
                   notes="This chart shows real values from our dataset next to their own Lag_1 (yesterday's "
                         "value) — notice they're often close, but not identical. That gap between 'today' and "
                         "'yesterday' is exactly the kind of thing a forecasting model tries to learn from.")

# ===========================================================================
# SLIDE 16 — WHAT IS A MOVING AVERAGE?
# ===========================================================================
slide_image_right(SEC7, "What Is a Moving Average?", "Slide 16",
                   ["“A moving average smooths short-term fluctuations so we can see the bigger pattern.”",
                    "A 7-day moving average replaces each day's value with the average of the last 7 days.",
                    "Original data: noisy, day-to-day.", "Moving average: a smoother, clearer pattern."],
                   "07_moving_average.png",
                   notes="Point directly at the chart: the grey line is the noisy original, the blue line is "
                         "its 7-day moving average — visibly calmer, and the underlying rhythm (including the "
                         "gentle upward trend) is much easier to see.")

# ===========================================================================
# SLIDE 17 — WHY USE MOVING AVERAGES?
# ===========================================================================
slide_bullets(SEC7, "Why Use Moving Averages?", "Slide 17",
              ["Reduce noise", "Identify trends", "Smooth fluctuations", "Understand overall behaviour",
               ("A moving average is useful for UNDERSTANDING the data — it does not automatically solve "
                "every forecasting problem.", 12, RED, True)],
              size=18,
              notes="This last bullet matters — a common beginner mistake is thinking 'compute a moving average' "
                    "IS the forecasting model. It's a descriptive/exploration tool first; we'll see in Hour 2 "
                    "that it can also become a FEATURE inside a real forecasting model, which is a more "
                    "sophisticated use of the same idea.")

# ===========================================================================
# SLIDE 18 — WHAT IS FORECASTING?
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC8, "What Is Forecasting?", "Slide 18")
add_text(s, LEFT_X, Inches(1.85), FULL_W, Inches(0.9),
          "“Forecasting means using historical data to estimate future values.”", size=20, color=INK, bold=True)
add_flow_horizontal(s, ["2024\nHistorical Data", "2025\n? ? ?\nForecast"], LEFT_X, Inches(3.3), FULL_W,
                     h=Inches(1.3), colors=[BLUE, AMBER])
add_notes(s, "Simple, deliberately bare-bones slide — the definition is the whole point. Everything from here "
             "forward in the course builds toward actually producing that '???' box with real numbers, using "
             "Python, in Hour 2.")

# ===========================================================================
# SLIDE 19 — SIMPLE FORECASTING BASELINE
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC8, "Simple Forecasting Baseline", "Slide 19")
add_box_with_text(s, LEFT_X, Inches(1.9), FULL_W, Inches(1.1), "“Tomorrow will be similar to today.”",
                   fill=BG_SOFT, text_color=INK, size=20, bold=True, line_color=LINE)
add_box_with_text(s, LEFT_X, Inches(3.2), FULL_W, Inches(1.1), "“Next week's value will be similar to last week's value.”",
                   fill=BG_SOFT, text_color=INK, size=20, bold=True, line_color=LINE)
add_text(s, LEFT_X, Inches(4.7), FULL_W, Inches(1.3),
          "This is called a NAIVE forecast — deliberately simple. A more advanced Machine Learning model must "
          "be COMPARED against something this simple to prove it's actually adding value.", size=17, color=SUBTEXT)
add_notes(s, "This is one of the most important practical habits in the entire course: always build the dumbest "
             "possible baseline FIRST. On our dataset, the naive forecast (today = tomorrow) has an MAE of "
             "about ₦8.2 million; we'll use that number directly in Hour 2 to judge whether our real model is "
             "actually worth the extra complexity.")

# ===========================================================================
# SLIDE 20 — THE FORECASTING WORKFLOW
# ===========================================================================
slide_flow_v(SEC8, "The Forecasting Workflow", "Slide 20",
             ["Collect Data", "Clean Data", "Visualize", "Understand Patterns", "Create Features",
              "Train Model", "Forecast", "Evaluate", "Business Decision"],
             box_w=Inches(5.6),
             notes="This is the master workflow slide — everything else in the course maps onto one of these "
                   "nine boxes. Consider returning to a simplified version of this diagram at the very end "
                   "(Slide 29) so learners see the full circle.")

# ===========================================================================
# SLIDE 21 — TIME SERIES DATA CANNOT BE TREATED LIKE ORDINARY DATA
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC9, "Time Series Data Cannot Be Shuffled", "Slide 21")
add_text(s, COL1_X, Inches(1.9), COL1_W, Inches(0.4), "INCORRECT", size=14, color=RED, bold=True)
add_box_with_text(s, COL1_X, Inches(2.3), COL1_W, Inches(2.3), "2024\n2022\n2025\n2023", fill=RGBColor(0xFE, 0xE2, 0xE2),
                   text_color=RED, size=20, bold=True, line_color=RED)
add_text(s, COL2_X, Inches(1.9), COL2_W, Inches(0.4), "CORRECT", size=14, color=TEAL, bold=True)
add_flow_vertical(s, ["Past Data", "Training Data", "Test Data", "Future"], COL2_X, Inches(2.3), COL2_W,
                   box_h=Inches(0.5), gap=Inches(0.18), fill=TEAL)
add_notes(s, "This directly attacks Misconception 1 head-on — some learners' instinct from ordinary ML is to "
             "randomly shuffle before train/test splitting (`train_test_split(shuffle=True)` is even the "
             "scikit-learn DEFAULT!). Say explicitly: for time series, that default is WRONG, and we'll show "
             "exactly how to split correctly in the next slide and in code.")

# ===========================================================================
# SLIDE 22 — TIME-BASED SPLITTING
# ===========================================================================
slide_image_full(SEC9, "Time-Based Splitting", "Slide 22", "08_train_test_split.png",
                  intro="We train using the past and test on the future — never the other way around.",
                  img_width=Inches(10.6), y=Inches(2.0),
                  notes="This is one of the most important beginner concepts in the whole course — call that "
                        "out explicitly. Point at the vertical dashed line: everything to the left trains the "
                        "model, everything to the right is held out to test it, and the split point respects "
                        "time — no test-period information leaks backward into training.")

# ===========================================================================
# SLIDE 23 — HOW DO WE KNOW OUR FORECAST IS GOOD?
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC10, "How Do We Know Our Forecast Is Good?", "Slide 23")
add_box_with_text(s, LEFT_X, Inches(1.85), Inches(3.75), Inches(2.5),
                   "MAE\n\n“On average, how far\naway were our\npredictions?”", fill=BLUE, size=16, radius=0.08)
add_box_with_text(s, Inches(4.78), Inches(1.85), Inches(3.75), Inches(2.5),
                   "RMSE\n\n“Gives more attention\nto large errors.”", fill=AMBER, size=16, radius=0.08)
add_box_with_text(s, Inches(8.55), Inches(1.85), Inches(4.0), Inches(2.5),
                   "MAPE\n\n“The average error as\na percentage of the\nactual value.”", fill=TEAL, size=16, radius=0.08)
add_text(s, LEFT_X, Inches(4.7), FULL_W, Inches(0.6), "We won't dwell on the formulas — the intuition matters far more.",
          size=15.5, color=MUTED, italic=True)
add_notes(s, "Keep the formulas off this slide entirely, as instructed — the plain-language framing IS the "
             "lesson. If asked for the formula: MAE = average of |actual - predicted|; RMSE = square root of the "
             "average of (actual - predicted)² (penalizes big misses more); MAPE = average of "
             "|actual - predicted| / actual, expressed as a percentage.")

# ===========================================================================
# SLIDE 24 — ACTUAL VS FORECAST
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC10, "Actual vs. Forecast", "Slide 24")
add_picture_framed(s, V + "09_actual_vs_forecast.png", Inches(0.78), Inches(1.75), w=Inches(7.5))
add_bullets(s, Inches(8.6), Inches(1.95), Inches(3.9), Inches(3.5),
            ["Naive baseline MAE: ₦8.2M", "Our model MAE: ₦4.7M",
             "The model tracks the weekly rhythm well but still misses some sharp spikes.",
             "The goal is not a perfect forecast — it's a USEFUL forecast with acceptable error "
             "for the business problem."], size=13.5)
add_notes(s, "These are real numbers computed on our own held-out test period in Hour 2 — not idealized "
             "textbook figures. Point out honestly where the dashed forecast line lags behind sharp actual "
             "spikes (holidays, unexplained events) — a good instructor moment to reinforce that 'noise' by "
             "definition can't be perfectly predicted.")

# ===========================================================================
# SLIDE 25 — WHERE DO WE START?
# ===========================================================================
slide_flow_v(SEC11, "Where Do We Start?", "Slide 25",
             ["1. Naive Forecast", "2. Moving Average", "3. Lag Features", "4. Simple Forecasting Model",
              "5. Advanced Models (later courses)"],
             colors=[MUTED, TEAL, BLUE, AMBER, RGBColor(0x9C, 0xA3, 0xAF)],
             box_w=Inches(6.2),
             notes="This progression is deliberately capped at step 4 for today — step 5 (ARIMA, SARIMA, "
                   "Prophet, LSTM) is explicitly out of scope, as stated on Slide 2. Reassure learners: steps "
                   "1-4 alone are genuinely useful and are what most real forecasting work actually starts with.")

# ===========================================================================
# SLIDE 26 — SIMPLE MODEL APPROACH
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC11, "Simple Model Approach", "Slide 26")
add_text(s, LEFT_X, Inches(1.85), FULL_W, Inches(0.6),
          "We can use previous values (and simple calendar facts) to predict future values.", size=17.5, color=INK, bold=True)
add_flow_vertical(s, ["Yesterday's Sales, Last Week's Sales,\nDay of Week, Month", "Prediction", "Tomorrow's Sales"],
                   Inches(3.8), Inches(2.8), Inches(5.7), box_h=Inches(0.85), gap=Inches(0.3), fill=BLUE)
add_text(s, LEFT_X, Inches(6.3), FULL_W, Inches(0.6),
          "This connects Time Series Forecasting directly to familiar Machine Learning: it's really just "
          "regression with the right features.", size=15, color=SUBTEXT, italic=True)
add_notes(s, "This is the big 'aha' moment for learners with an existing ML background: forecasting doesn't "
             "require an entirely new toolkit. Once you've built lag and calendar features, you can hand the "
             "problem to a completely ordinary regression model — which is exactly what we do in Notebook 03.")

# ===========================================================================
# SLIDE 27 — FORECASTING IS NOT JUST ABOUT THE MODEL
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC12, "Forecasting Is Not Just About the Model", "Slide 27")
add_flow_vertical(s, ["Forecast", "Business Interpretation", "Action"], Inches(4.0), Inches(1.85), Inches(5.3),
                   box_h=Inches(0.65), gap=Inches(0.28), fill=TEAL)
add_text(s, LEFT_X, Inches(4.65), FULL_W, Inches(0.5), "Example: predicted ATM withdrawals increase.", size=16, color=MUTED, bold=True)
add_bullets(s, LEFT_X, Inches(5.15), FULL_W, Inches(1.6),
            ["Supply more cash to affected branches/ATMs", "Improve ATM cash-loading planning",
             "Reduce the risk of cash shortages during peak demand"], size=16)
add_notes(s, "This is the payoff slide — the forecast only matters once it changes what the bank actually DOES. "
             "Ask: 'what's the cost of under-forecasting cash demand? What's the cost of over-forecasting it?' "
             "(Under: angry customers, reputational damage. Over: idle cash, opportunity cost, security risk.) "
             "Both errors are real costs — a good forecast helps balance them.")

# ===========================================================================
# SLIDE 28 — COMMON MISTAKES
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC12, "Common Mistakes", "Slide 28")
add_bullets(s, LEFT_X, Inches(1.85), FULL_W, Inches(3.3),
            ["Randomly shuffling Time Series data", "Ignoring missing dates",
             "Ignoring seasonality", "Using future information to predict the past",
             "Not comparing against a baseline", "Looking only at model accuracy",
             "Forgetting business context"], size=17)
add_pill(s, Inches(3.0), Inches(5.5), Inches(7.3), Inches(0.65), "The 4th one has a name: DATA LEAKAGE",
          fill=RED, text_color=WHITE, size=16)
add_notes(s, "Spend extra time on data leakage — it's the single most damaging and most common beginner "
             "mistake in time series work, and it can make a model look GREAT in testing while being useless in "
             "production. Simple example: accidentally including 'this month's total sales' as a feature to "
             "predict 'today's sales' — at prediction time, in the real world, you would NOT yet know the "
             "month's total (it includes days that haven't happened yet).")

# ===========================================================================
# SLIDE 29 — FINAL RECAP
# ===========================================================================
slide_flow_v(SEC12, "Final Recap", "Slide 29",
             ["Time Series", "Time Order Matters", "Visualize", "Find Patterns (Trend, Seasonality, Noise)",
              "Create Lags", "Train on the Past", "Test on the Future", "Forecast", "Evaluate", "Business Action"],
             box_w=Inches(5.6),
             notes="This is the single-slide summary of the ENTIRE conceptual hour — consider lingering here "
                   "for questions before the coding break. Every slide in Hour 1 maps to one box in this chain.")

# ===========================================================================
# SLIDE 30 — TRANSITION TO CODING
# ===========================================================================
slide_quote(SEC12, "Transition to Coding", "Slide 30",
            "“Now that we understand how forecasting works,\nhow do we build one in Python?”",
            subtext="Let's open the notebooks.",
            notes="Short, energetic transition slide. Take a 2-3 minute break here if the room needs it before "
                  "diving into Hour 2's hands-on notebooks.")

# ===========================================================================
# SLIDE 31 — IMPORTANT DISTINCTIONS
# ===========================================================================
slide_table("Instructor Reference", "Important Distinctions", "Slide 31",
            ["Myth", "Reality"],
            [["Time Series data can be randomly shuffled.", "The order of observations matters — shuffling "
              "destroys the time dependency we're trying to model."],
             ["A good ML model automatically produces a good forecast.", "Forecasting depends on data quality, "
              "patterns, features, model choice, validation, and business context."],
             ["More complex models are always better.", "A simple baseline may perform surprisingly well — "
              "always compare against one."],
             ["Forecasting means predicting perfectly.", "Forecasting means estimating the future WITH uncertainty."],
             ["High accuracy is the only thing that matters.", "The forecast must also be USEFUL for the "
              "business decision at hand."],
             ["Future information can be used as a feature.", "Using information not yet available at prediction "
              "time causes DATA LEAKAGE."]],
            col_widths=[1, 1.5], table_h=Inches(5.0),
            notes="Consider running this as a rapid true/false check with the class before the quiz — these six "
                  "distinctions are the ones beginners misstate most often in follow-up conversations.")

# ===========================================================================
# SLIDE 32 — QUICK QUIZ
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, "Knowledge Check", "Quick Quiz", "Slide 32")
quiz = ["1. What is Time Series data?", "2. Why is randomly shuffling Time Series data usually a problem?",
        "3. What is a trend?", "4. What is seasonality?", "5. What is a lag?",
        "6. Why should we create a baseline forecast?", "7. What does MAE measure?",
        "8. What is data leakage in Time Series Forecasting?"]
add_bullets(s, COL1_X, CONTENT_TOP, COL1_W, Inches(5.3), quiz[:4], size=17, space_after=20, marker="")
add_bullets(s, COL2_X, CONTENT_TOP, COL2_W, Inches(5.3), quiz[4:], size=17, space_after=20, marker="")
add_notes(s, "ANSWER KEY (for instructor use):\n"
             "1. A sequence of observations recorded over time, where the order of observations matters.\n"
             "2. Because it destroys the time dependency in the data — training and testing must respect the "
             "natural order (train on the past, test on the future), otherwise the model could 'see' future "
             "information during training that would not be available at real prediction time.\n"
             "3. The general direction of the data over a long period of time (upward, downward, or none).\n"
             "4. A pattern that repeats at regular intervals (e.g. every week, every month, every year).\n"
             "5. A previous value of a variable (e.g. Lag_1 = yesterday's value, Lag_7 = the value from 7 days ago).\n"
             "6. So we have something simple to compare a more advanced model against — if the advanced model "
             "can't beat the baseline, its extra complexity isn't paying off.\n"
             "7. Mean Absolute Error — on average, how far away were our predictions from the actual values.\n"
             "8. Using information that would not actually have been available at the time of prediction as a "
             "feature — this makes a model look artificially accurate during testing but fail in the real world.")

# ===========================================================================
# SLIDE 33 — REFERENCES / SOURCES
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, "Appendix", "References / Sources", "Slide 33")
add_text(s, LEFT_X, CONTENT_TOP, FULL_W, Inches(0.32), "Textbook", size=14, color=MUTED, bold=True)
add_bullets(s, LEFT_X, Inches(1.92), FULL_W, Inches(1.15),
            ["Brockwell, P. J., & Davis, R. A. (2016). Introduction to Time Series and Forecasting (3rd ed.). "
             "Springer Texts in Statistics. — Chapter 1, foundational examples of trend, seasonality, and "
             "moving-average smoothing."],
            size=14, space_after=6)
add_text(s, LEFT_X, Inches(3.05), FULL_W, Inches(0.32), "Official Documentation (verified for current syntax)",
          size=14, color=MUTED, bold=True)
add_bullets(s, LEFT_X, Inches(3.42), FULL_W, Inches(2.7),
            ["pandas API Reference — DataFrame.rolling(), DataFrame.shift() "
             "(pandas.pydata.org/docs/reference/ ...)",
             "scikit-learn API Reference — mean_absolute_error, root_mean_squared_error, "
             "mean_absolute_percentage_error, LinearRegression "
             "(scikit-learn.org/stable/modules/generated/ ...)",
             "scikit-learn User Guide — Regression metrics "
             "(scikit-learn.org/stable/modules/model_evaluation.html)",
             "Every code snippet in this deck and the notebooks was checked against these pages during "
             "development."],
            size=13, space_after=6)
add_text(s, LEFT_X, Inches(6.35), FULL_W, Inches(0.32), "Data", size=14, color=MUTED, bold=True)
add_bullets(s, LEFT_X, Inches(6.72), FULL_W, Inches(0.6),
            ["All withdrawal data used in this course is SYNTHETIC (data/generate_data.py). It does not "
             "represent any real bank, customer, or transaction."], size=12.5, space_after=4)
add_notes(s, "All pandas and scikit-learn syntax in this deck (rolling, shift, LinearRegression, the three "
             "regression metrics) was checked against the official documentation during course development to "
             "ensure it reflects current, non-deprecated APIs — including root_mean_squared_error, added in "
             "scikit-learn 1.4 as the modern replacement for mean_squared_error(squared=False).")

# ===========================================================================
# Final pass: footers
# ===========================================================================
total = len(prs.slides)
for i, sl in enumerate(prs.slides, start=1):
    add_footer(sl, i, total, tag="Time Series Forecasting")

prs.save("Time_Series_Forecasting_Beginner_Course.pptx")
print(f"Saved deck with {total} slides.")
