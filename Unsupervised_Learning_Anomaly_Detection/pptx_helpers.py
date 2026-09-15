"""
Reusable layout helpers for building the Unsupervised Learning (K-Means & PCA)
course deck with python-pptx. Keeping these in one module keeps build_pptx.py
focused on CONTENT rather than layout plumbing.

Design system:
  - 16:9 canvas, clean light background
  - Dark slate text on off-white background
  - Blue / amber / teal / purple accent palette (matches the chart palette
    used in visuals/generate_visuals.py so slides and charts feel like one system)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
import copy

# ---------------------------------------------------------------- palette --
INK = RGBColor(0x1F, 0x29, 0x37)
SUBTEXT = RGBColor(0x4B, 0x55, 0x63)
MUTED = RGBColor(0x6B, 0x72, 0x80)
BG = RGBColor(0xFF, 0xFF, 0xFF)
BG_SOFT = RGBColor(0xF7, 0xF9, 0xFB)
LINE = RGBColor(0xE5, 0xE7, 0xEB)

BLUE = RGBColor(0x25, 0x63, 0xEB)
BLUE_DARK = RGBColor(0x1E, 0x40, 0xAF)
BLUE_SOFT = RGBColor(0xDB, 0xEA, 0xFE)
AMBER = RGBColor(0xD9, 0x77, 0x06)
AMBER_SOFT = RGBColor(0xFD, 0xE9, 0xC8)
TEAL = RGBColor(0x0D, 0x94, 0x88)
TEAL_SOFT = RGBColor(0xCC, 0xEC, 0xE9)
PURPLE = RGBColor(0x7C, 0x3A, 0xED)
PURPLE_SOFT = RGBColor(0xE9, 0xDD, 0xFC)
RED = RGBColor(0xDC, 0x26, 0x26)
DARK_PANEL = RGBColor(0x11, 0x18, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"
FONT_MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

CLUSTER_COLORS = [BLUE, AMBER, TEAL, PURPLE]


def new_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs, bg=BG):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return slide


def set_font(run, size=18, color=INK, bold=False, italic=False, font=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.0,
             word_wrap=True, shrink=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        set_font(run, size, color, bold, italic, font)
    return box


def add_bullets(slide, x, y, w, h, items, size=16, color=SUBTEXT, bold_first=False,
                 space_after=10, font=FONT, marker="•  ", anchor=MSO_ANCHOR.TOP):
    """items: list of str, or list of (str, size, color, bold) tuples for mixed styling,
    or list of (text, level) for indentation."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        level = 0
        txt = item
        item_size, item_color, item_bold = size, color, False
        if isinstance(item, tuple):
            if len(item) == 2:
                txt, level = item
            elif len(item) == 4:
                txt, item_size, item_color, item_bold = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.line_spacing = 1.12
        prefix = ("      " * level) + (marker if marker else "")
        run = p.add_run()
        run.text = f"{prefix}{txt}"
        set_font(run, item_size, item_color, item_bold, False, font)
    return box


def add_rect(slide, x, y, w, h, fill=BLUE, line_color=None, radius=None, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    if radius:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp


def add_box_with_text(slide, x, y, w, h, text, fill=BLUE, text_color=WHITE, size=15,
                       bold=True, radius=0.12, align=PP_ALIGN.CENTER, line_color=None,
                       font=FONT):
    shp = add_rect(slide, x, y, w, h, fill=fill, radius=radius, line_color=line_color)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        set_font(run, size, text_color, bold, False, font)
    return shp


def add_kicker_title(slide, kicker, title, section=None, title_size=30):
    """Top-of-slide heading: small accent bar + kicker label + big title, optional
    section tag in the top-right corner."""
    add_rect(slide, Inches(0.55), Inches(0.42), Inches(0.09), Inches(0.62), fill=BLUE)
    if kicker:
        add_text(slide, Inches(0.78), Inches(0.40), Inches(9.5), Inches(0.3), kicker.upper(),
                  size=12.5, color=BLUE, bold=True)
        add_text(slide, Inches(0.78), Inches(0.66), Inches(11.8), Inches(0.62), title,
                  size=title_size, color=INK, bold=True)
    else:
        add_text(slide, Inches(0.78), Inches(0.40), Inches(11.8), Inches(0.88), title,
                  size=title_size, color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if section:
        add_text(slide, Inches(10.6), Inches(0.42), Inches(2.2), Inches(0.3), section.upper(),
                  size=10.5, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)
    add_rect(slide, Inches(0.55), Inches(1.28), Inches(12.23), Pt(1.4), fill=LINE)


def add_footer(slide, page_no, total, tag="Unsupervised Learning · K-Means & PCA"):
    add_text(slide, Inches(0.55), Inches(7.14), Inches(6), Inches(0.28), tag, size=9.5, color=MUTED)
    add_text(slide, Inches(11.8), Inches(7.14), Inches(1.0), Inches(0.28), f"{page_no}", size=9.5,
              color=MUTED, align=PP_ALIGN.RIGHT)


def add_picture_framed(slide, path, x, y, w=None, h=None, caption=None, border=True):
    pic = slide.shapes.add_picture(path, x, y, width=w, height=h)
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(1)
    if caption:
        add_text(slide, x, y + pic.height + Inches(0.06), pic.width, Inches(0.32), caption,
                  size=11.5, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
    return pic


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_arrow_between(slide, x1, y1, x2, y2, color=MUTED, weight=2.25):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(weight)
    line = conn.line._get_or_add_ln()
    tail = line.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    line.append(tail)
    return conn


def add_flow_vertical(slide, steps, x, y, w, box_h=Inches(0.62), gap=Inches(0.30),
                       fill=BLUE, text_color=WHITE, size=15):
    cy = y
    for i, step in enumerate(steps):
        add_box_with_text(slide, x, cy, w, box_h, step, fill=fill, text_color=text_color, size=size, radius=0.18)
        cy = Emu(cy + box_h + gap)
        if i < len(steps) - 1:
            cx = Emu(x + w // 2)
            add_arrow_between(slide, cx, Emu(cy - gap + Pt(2)), cx, Emu(cy - Pt(2)), color=MUTED, weight=2.25)
    return cy


def add_flow_horizontal(slide, steps, x, y, w, h=Inches(1.0), gap=Inches(0.35),
                         fill=BLUE, text_color=WHITE, size=14.5, colors=None):
    n = len(steps)
    box_w = Emu(int((w - gap * (n - 1)) / n))
    cx = x
    for i, step in enumerate(steps):
        c = colors[i % len(colors)] if colors else fill
        add_box_with_text(slide, cx, y, box_w, h, step, fill=c, text_color=text_color, size=size, radius=0.16)
        nx = Emu(cx + box_w + gap)
        if i < n - 1:
            ay = Emu(y + h // 2)
            add_arrow_between(slide, Emu(cx + box_w + Pt(2)), ay, Emu(nx - Pt(2)), ay, color=MUTED, weight=2.25)
        cx = nx
    return box_w


def add_table(slide, x, y, w, h, headers, rows, header_fill=INK, header_color=WHITE,
              body_size=14, header_size=14.5, col_widths=None, zebra=True,
              highlight_col=None, highlight_fill=BLUE_SOFT):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gshape = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    table = gshape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Emu(int(w * cw / total))
    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = htext
        set_font(run, header_size, header_color, True)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Pt(10); cell.margin_right = Pt(10)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(val)
            set_font(run, body_size, INK if c == 0 else SUBTEXT, c == 0)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Pt(10); cell.margin_right = Pt(10)
            if highlight_col is not None and c == highlight_col:
                cell.fill.solid()
                cell.fill.fore_color.rgb = highlight_fill
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = BG_SOFT if (zebra and r % 2 == 0) else BG
    # remove default banding style visuals by keeping our own fills (leave table style as-is; explicit fills override)
    return table


def add_code_block(slide, x, y, w, h, code, size=14.5, fill=DARK_PANEL,
                    text_color=RGBColor(0xE5, 0xE7, 0xEB), keyword_color=RGBColor(0x7D, 0xD3, 0xFC)):
    shp = add_rect(slide, x, y, w, h, fill=fill, radius=0.06)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(16); tf.margin_right = Pt(16); tf.margin_top = Pt(10); tf.margin_bottom = Pt(10)
    lines = code.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.25
        run = p.add_run()
        run.text = line if line.strip() != "" else " "
        set_font(run, size, text_color, False, False, FONT_MONO)
    return shp


def add_pill(slide, x, y, w, h, text, fill=BLUE_SOFT, text_color=BLUE_DARK, size=12):
    return add_box_with_text(slide, x, y, w, h, text, fill=fill, text_color=text_color, size=size, bold=True, radius=0.5)
