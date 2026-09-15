# -*- coding: utf-8 -*-
"""
Builds Collaborative_Filtering_Engines_Beginner_Course.pptx

Run from the project root:
    python build_pptx.py
"""

from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx_helpers import *

V = "visuals/"

prs = new_presentation()

CONTENT_TOP = Inches(1.55)
LEFT_X = Inches(0.78)
FULL_W = Inches(11.78)
COL1_X = Inches(0.78)
COL1_W = Inches(5.55)
COL2_X = Inches(6.60)
COL2_W = Inches(6.00)


def slide_bullets(kicker, title, section, bullets, notes="", subtitle=None, size=17.5, marker="•  "):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    top = CONTENT_TOP
    if subtitle:
        add_text(s, LEFT_X, top, FULL_W, Inches(0.5), subtitle, size=16, color=MUTED, italic=True)
        top = Emu(top + Inches(0.55))
    add_bullets(s, LEFT_X, top, FULL_W, Inches(6.95) - top, bullets, size=size, marker=marker)
    if notes:
        add_notes(s, notes)
    return s


def slide_image_right(kicker, title, section, bullets, image, notes="", caption=None,
                       img_width=Inches(5.85), subtitle=None):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    top = CONTENT_TOP
    if subtitle:
        add_text(s, COL1_X, top, COL1_W, Inches(0.5), subtitle, size=15.5, color=MUTED, italic=True)
        top = Emu(top + Inches(0.55))
    add_bullets(s, COL1_X, top, COL1_W, Inches(6.9) - top, bullets, size=16.5)
    add_picture_framed(s, V + image, COL2_X, Inches(1.85), w=img_width, caption=caption)
    if notes:
        add_notes(s, notes)
    return s


def slide_image_full(kicker, title, section, image, notes="", caption=None,
                      img_width=Inches(9.4), intro=None, y=Inches(1.75)):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    if intro:
        add_text(s, LEFT_X, CONTENT_TOP, FULL_W, Inches(0.5), intro, size=16.5, color=MUTED, italic=True)
        y = Emu(y + Inches(0.35))
    x = Emu(int((SLIDE_W - img_width) / 2))
    add_picture_framed(s, V + image, x, y, w=img_width, caption=caption)
    if notes:
        add_notes(s, notes)
    return s


def slide_flow_v(kicker, title, section, steps, notes="", intro=None, colors=None, box_w=Inches(5.6)):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    y = CONTENT_TOP
    if intro:
        add_text(s, LEFT_X, y, FULL_W, Inches(0.5), intro, size=16.5, color=MUTED, italic=True)
        y = Emu(y + Inches(0.5))
    x = Emu(int((SLIDE_W - box_w) / 2))
    n = len(steps)
    if n > 8:
        box_h, gap, fsize = Inches(0.44), Inches(0.09), 13
    elif n > 6:
        box_h, gap, fsize = Inches(0.5), Inches(0.16), 14.5
    elif n > 5:
        box_h, gap, fsize = Inches(0.56), Inches(0.20), 15
    else:
        box_h, gap, fsize = Inches(0.62), Inches(0.26), 15
    if colors is None:
        cy = y
        for i, step in enumerate(steps):
            add_box_with_text(s, x, cy, box_w, box_h, step, fill=BLUE, radius=0.18, size=fsize)
            cy = Emu(cy + box_h + gap)
            if i < n - 1:
                cx = Emu(x + box_w // 2)
                add_arrow_between(s, cx, Emu(cy - gap + Pt(2)), cx, Emu(cy - Pt(2)))
    else:
        cy = y
        for i, (step, c) in enumerate(zip(steps, colors)):
            add_box_with_text(s, x, cy, box_w, box_h, step, fill=c, radius=0.18, size=fsize)
            cy = Emu(cy + box_h + gap)
            if i < n - 1:
                cx = Emu(x + box_w // 2)
                add_arrow_between(s, cx, Emu(cy - gap + Pt(2)), cx, Emu(cy - Pt(2)))
    if notes:
        add_notes(s, notes)
    return s


def slide_flow_h(kicker, title, section, steps, notes="", intro=None, colors=None, y=Inches(3.1), h=Inches(1.0)):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    if intro:
        add_text(s, LEFT_X, CONTENT_TOP, FULL_W, Inches(0.5), intro, size=16.5, color=MUTED, italic=True)
    add_flow_horizontal(s, steps, LEFT_X, y, FULL_W, h=h, colors=colors)
    if notes:
        add_notes(s, notes)
    return s


def slide_table(kicker, title, section, headers, rows, notes="", intro=None,
                 col_widths=None, table_h=Inches(3.6), header_fill=INK):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    top = CONTENT_TOP
    if intro:
        add_text(s, LEFT_X, top, FULL_W, Inches(0.5), intro, size=16, color=MUTED, italic=True)
        top = Emu(top + Inches(0.55))
    add_table(s, LEFT_X, top, FULL_W, table_h, headers, rows, col_widths=col_widths, header_fill=header_fill)
    if notes:
        add_notes(s, notes)
    return s


def slide_code(kicker, title, section, code, bullets=None, notes="", code_h=Inches(2.2), intro=None):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    top = CONTENT_TOP
    if intro:
        add_text(s, LEFT_X, top, FULL_W, Inches(0.45), intro, size=16, color=MUTED, italic=True)
        top = Emu(top + Inches(0.5))
    add_code_block(s, LEFT_X, top, FULL_W, code_h, code)
    if bullets:
        bullet_top = Emu(top + code_h + Inches(0.3))
        bullet_h = Emu(int(Inches(6.9)) - int(bullet_top))
        add_bullets(s, LEFT_X, bullet_top, FULL_W, bullet_h, bullets, size=15.5)
    if notes:
        add_notes(s, notes)
    return s


def slide_quote(kicker, title, section, quote, subtext=None, notes="", quote_size=32, fill=DARK_PANEL,
                 text_color=WHITE):
    s = blank_slide(prs)
    add_kicker_title(s, kicker, title, section=section)
    panel = add_rect(s, LEFT_X, Inches(2.1), FULL_W, Inches(3.6), fill=fill, radius=0.06)
    tf = panel.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(30); tf.margin_right = Pt(30)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = quote
    set_font(run, quote_size, text_color, True)
    if subtext:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(14)
        run2 = p2.add_run()
        run2.text = subtext
        set_font(run2, 17, RGBColor(0xC7, 0xCE, 0xD9), False, True)
    if notes:
        add_notes(s, notes)
    return s


SEC1 = "Section 1 · Introduction"
SEC2 = "Section 2 · The Real Problem"
SEC3 = "Section 3 · What Is a Recommender?"
SEC4 = "Section 4 · User-Item Interactions"
SEC5 = "Section 5 · Feedback Types"
SEC6 = "Section 6 · The User-Item Matrix"
SEC7 = "Section 7 · Collaborative Filtering"
SEC8 = "Section 8 · User-Based Filtering"
SEC9 = "Section 9 · Item-Based Filtering"
SEC10 = "Section 10 · User vs. Item"
SEC11 = "Section 11 · Similarity"
SEC12 = "Section 12 · Building the Engine"
SEC13 = "Section 13 · Making Recommendations"
SEC14 = "Section 14 · Evaluation"
SEC15 = "Section 15 · Real-World Challenges"
SEC16 = "Section 16 · Misconceptions"
SEC17 = "Section 17 · Business Value"
SEC18 = "Section 18 · Recap"

# ===========================================================================
# SLIDE 1 — TITLE
# ===========================================================================
s = blank_slide(prs, bg=DARK_PANEL)
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(0.14), fill=PURPLE)
add_text(s, Inches(0.9), Inches(1.1), Inches(7), Inches(0.4), "BEGINNER DATA SCIENCE COURSE",
          size=14, color=RGBColor(0xDD, 0xD6, 0xFE), bold=True)
add_text(s, Inches(0.9), Inches(2.2), Inches(11.6), Inches(1.3), "Building Collaborative Filtering Engines",
          size=40, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(3.15), Inches(11.5), Inches(0.9),
          "Recommending What Users May Like Based on Similar Behaviour", size=22,
          color=RGBColor(0xDD, 0xD6, 0xFE), bold=True)
add_rect(s, Inches(0.9), Inches(4.15), Inches(1.3), Pt(3.5), fill=AMBER)
add_text(s, Inches(0.9), Inches(4.35), Inches(10.5), Inches(0.6),
          "Time Series & Recommender Systems", size=18, color=RGBColor(0xE5, 0xE7, 0xEB), italic=True)
for i, (label, c) in enumerate([("Users", BLUE), ("Preferences", TEAL), ("Patterns", AMBER), ("Recommendations", PURPLE)]):
    add_pill(s, Inches(0.9 + i * 2.85), Inches(5.65), Inches(2.6), Inches(0.5), label, fill=c, text_color=WHITE, size=13)
add_notes(s, "Welcome the class. This is a 2-hour session: Hour 1 is concepts (this deck), Hour 2 is hands-on "
             "Python. Frame it as building ONE recommendation engine together for a movie streaming platform, "
             "not a tour of separate algorithms. Ask: has anyone ever wondered how Netflix or Spotify seems to "
             "'know' what they might like? That's exactly what we're demystifying today — no magic, just "
             "patterns in what similar people like.")

# ===========================================================================
# SLIDE 2 — LEARNING OBJECTIVES
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, "Introduction", "Learning Objectives", section="Slide 2")
add_text(s, LEFT_X, CONTENT_TOP, FULL_W, Inches(0.5), "By the end of this session, you will be able to:",
          size=16, color=MUTED, italic=True)
objs = ["Explain what a recommendation system is.", "Understand user-item interactions.",
        "Explain explicit and implicit feedback.", "Understand collaborative filtering.",
        "Explain user-based collaborative filtering.", "Explain item-based collaborative filtering.",
        "Understand similarity.", "Build a simple recommendation engine in Python.",
        "Generate recommendations.", "Evaluate basic recommendation results."]
add_bullets(s, LEFT_X, Inches(2.1), COL1_W, Inches(4.5), objs[:5], size=17, space_after=14)
add_bullets(s, COL2_X, Inches(2.1), COL2_W, Inches(4.5), objs[5:], size=17, space_after=14)
add_notes(s, "Walk down the list quickly — this is a map, not a lesson. Mention explicitly that we are "
             "deliberately NOT covering matrix factorization, SVD, ALS, or neural collaborative filtering in "
             "depth today — those are later-course territory. Today is entirely about building strong intuition "
             "for neighborhood-based collaborative filtering, plus one working, understandable engine.")

# ===========================================================================
# SLIDE 3 — WHY DO WE NEED RECOMMENDATION SYSTEMS?
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC2, "Why Do We Need Recommendation Systems?", section="Slide 3")
examples = ["Netflix recommending movies", "Spotify recommending songs", "Amazon recommending products",
            "YouTube recommending videos", "Online stores recommending products"]
for i, ex in enumerate(examples):
    x = LEFT_X + Inches((i % 3) * 3.95)
    y = Inches(1.9) + Inches((i // 3) * 0.85)
    add_pill(s, x, y, Inches(3.7), Inches(0.6), ex, fill=PURPLE_SOFT, text_color=RGBColor(0x5B, 0x21, 0xB6), size=14)
add_text(s, LEFT_X, Inches(3.9), FULL_W, Inches(0.6), "How does Netflix know what movie you may like?",
          size=20, color=INK, bold=True)
add_text(s, LEFT_X, Inches(4.7), FULL_W, Inches(1.0),
          "It does not always know exactly. It learns from patterns in user behaviour.", size=18, color=SUBTEXT, italic=True)
add_notes(s, "Ask learners to name a platform's recommendation feature they've personally noticed working well "
             "(or badly!) before revealing these examples. The key reframe on this slide: these platforms are "
             "not reading minds — they're pattern-matching on behaviour, which is exactly what we'll build by "
             "hand in Hour 2.")

# ===========================================================================
# SLIDE 4 — THE MOVIE STREAMING PROBLEM
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC2, "The Movie Streaming Problem", section="Slide 4")
add_box_with_text(s, LEFT_X, Inches(1.9), Inches(3.75), Inches(1.7), "1,000,000\nusers", fill=BLUE, size=22, radius=0.08)
add_box_with_text(s, Inches(4.78), Inches(1.9), Inches(3.75), Inches(1.7), "50,000\nmovies", fill=TEAL, size=22, radius=0.08)
add_box_with_text(s, Inches(8.55), Inches(1.9), Inches(4.0), Inches(1.7), "Millions of\ninteractions", fill=AMBER, size=22, radius=0.08)
add_text(s, LEFT_X, Inches(4.1), FULL_W, Inches(0.6), "How can we decide what movie to recommend to each user?",
          size=19, color=INK, bold=True)
add_pill(s, Inches(2.6), Inches(5.2), Inches(7.6), Inches(0.65), "Showing every movie to every user is not practical",
          fill=RED, text_color=WHITE, size=15.5)
add_notes(s, "Make the scale concrete — a million users and 50,000 movies means personalization isn't a nice-to-"
             "have, it's the only way the platform stays usable. Ask: 'if Netflix just showed everyone the same "
             "homepage, what would happen?' (Answer: most users would never find anything relevant, and give "
             "up browsing.)")

# ===========================================================================
# SLIDE 5 — SIMPLE DEFINITION
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC3, "What Is a Recommender System?", "Slide 5")
add_text(s, LEFT_X, Inches(1.9), FULL_W, Inches(0.9),
          "“A recommender system helps users discover items\nthey may be interested in.”", size=22, color=INK, bold=True)
items = ["Movies", "Songs", "Products", "Books", "Courses", "Jobs", "News Articles"]
for i, it in enumerate(items):
    x = LEFT_X + Inches((i % 4) * 2.95)
    y = Inches(3.6) + Inches((i // 4) * 0.75)
    add_pill(s, x, y, Inches(2.75), Inches(0.55), it, fill=BLUE_SOFT, text_color=BLUE_DARK, size=14)
add_notes(s, "Emphasize the word 'ITEM' as the general umbrella term — it applies to anything a platform can "
             "recommend, not just movies. We'll use movies as our running example for the rest of the course "
             "because it's the most intuitive, but the exact same technique applies to any of these categories.")

# ===========================================================================
# SLIDE 6 — THE BASIC IDEA
# ===========================================================================
slide_flow_v(SEC3, "The Basic Idea", "Slide 6",
             ["User", "Past Behaviour", "Find Patterns", "Predict Interest", "Recommend Item"],
             box_w=Inches(5.4),
             notes="Recommendations are PREDICTIONS about possible user preferences — not certainties. Plant "
                   "that seed early; it directly pre-empts Misconception 3 later (a recommendation does not "
                   "guarantee the user will like the item).")

# ===========================================================================
# SLIDE 7 — USERS AND ITEMS
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC4, "Users and Items", "Slide 7")
add_box_with_text(s, LEFT_X, Inches(2.0), Inches(5.6), Inches(2.2), "USER\n\nA person interacting\nwith the system.\n\ne.g. John", fill=BLUE, size=17, radius=0.08)
add_box_with_text(s, Inches(6.85), Inches(2.0), Inches(5.6), Inches(2.2),
                   "ITEM\n\nSomething the system\ncan recommend.\n\ne.g. Movie A, B, C, D", fill=TEAL, size=17, radius=0.08)
add_notes(s, "Simple vocabulary-setting slide — keep it brief. These two words (USER, ITEM) will appear "
             "constantly for the rest of the course, so make sure everyone's comfortable with them before "
             "moving on.")

# ===========================================================================
# SLIDE 8 — WHAT IS AN INTERACTION?
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC4, "What Is an Interaction?", "Slide 8")
add_flow_horizontal(s, ["User", "Interaction", "Item"], LEFT_X, Inches(1.9), FULL_W, h=Inches(1.0),
                     colors=[BLUE, AMBER, TEAL])
add_text(s, LEFT_X, Inches(3.3), FULL_W, Inches(0.4), "Examples of interactions:", size=15, color=MUTED, bold=True)
add_bullets(s, LEFT_X, Inches(3.75), FULL_W, Inches(2.5),
            ["Watching a movie", "Rating a movie", "Buying a product", "Clicking a video",
             "Listening to a song", "Adding a product to a cart"], size=16.5)
add_notes(s, "Users interact with items in many different ways — this variety matters because it directly sets "
             "up the next two slides (explicit vs. implicit feedback). Ask: 'which of these examples tells us "
             "MOST clearly what the user thinks — and which tells us the LEAST?'")

# ===========================================================================
# SLIDE 9 — EXPLICIT FEEDBACK
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC5, "Explicit Feedback", "Slide 9")
add_text(s, LEFT_X, Inches(1.9), FULL_W, Inches(0.8),
          "“Explicit feedback happens when a user directly\ntells us what they think.”", size=20, color=INK, bold=True)
add_bullets(s, LEFT_X, Inches(3.2), FULL_W, Inches(1.4),
            ["Movie rating: 5 stars", "Product rating: 4 stars", "Like / dislike"], size=17)
add_box_with_text(s, Inches(3.4), Inches(4.7), Inches(6.9), Inches(0.9), "John  →  Movie A  →  ★★★★★",
                   fill=BG_SOFT, text_color=INK, size=19, line_color=LINE)
add_notes(s, "This is the clearest, most direct kind of signal — the user is TELLING us their preference. Our "
             "movie ratings dataset in Hour 2 is entirely explicit feedback (1-5 star ratings), which is why "
             "it's a good first dataset for beginners to learn on.")

# ===========================================================================
# SLIDE 10 — IMPLICIT FEEDBACK
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC5, "Implicit Feedback", "Slide 10")
add_text(s, LEFT_X, Inches(1.9), FULL_W, Inches(0.6), "“Implicit feedback comes from user behaviour.”",
          size=20, color=INK, bold=True)
add_bullets(s, LEFT_X, Inches(2.75), FULL_W, Inches(1.8),
            ["Clicks", "Purchases", "Watch history", "Listening history", "Search history",
             "Time spent on a product"], size=16)
add_flow_vertical(s, ["User watched Movie A", "User watched Movie B", "User watched Movie C"],
                   Inches(4.0), Inches(4.85), Inches(5.3), box_h=Inches(0.45), gap=Inches(0.12), fill=AMBER)
add_notes(s, "The critical distinction: 'the user did not directly SAY \"I like this\" — but their behaviour "
             "gives us useful clues.' Most real-world recommendation data is actually implicit (few people "
             "bother to rate things, but everyone leaves behavioural traces) — flag that our course dataset "
             "uses explicit ratings for teaching simplicity, but production systems often lean heavily on "
             "implicit signals instead.")

# ===========================================================================
# SLIDE 11 — TURN BEHAVIOUR INTO DATA
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC6, "Turn Behaviour Into Data", "Slide 11")
add_text(s, LEFT_X, Inches(1.65), FULL_W, Inches(0.4), "The User-Item Matrix", size=18, color=BLUE, bold=True)
add_table(s, LEFT_X, Inches(2.1), Inches(7.0), Inches(2.1),
          ["User", "Movie A", "Movie B", "Movie C", "Movie D"],
          [["John", "5", "4", "?", "2"], ["Mary", "5", "4", "5", "1"], ["David", "1", "2", "4", "5"]],
          col_widths=[1, 0.8, 0.8, 0.8, 0.8], header_fill=INK)
add_bullets(s, Inches(8.2), Inches(2.15), Inches(3.9), Inches(3.0),
            ["Rows = Users", "Columns = Items", "Values = Ratings or interactions",
             "Missing values = Unknown preferences"], size=14.5)
add_notes(s, "This is one of the anchor slides for the whole course — the user-item matrix is the single data "
             "structure that everything else builds on. Point at the '?' for John/Movie C explicitly: that gap "
             "is exactly what we're trying to fill in. This same table appears again on the very next slide.")

# ===========================================================================
# SLIDE 12 — THE BIG QUESTION
# ===========================================================================
slide_quote(SEC6, "The Big Question", "Slide 12",
            "“How can we fill in the missing preferences?”",
            subtext="For example: John + Movie C = ?  —  this becomes the recommendation problem.",
            notes="Let this sit for a moment — it's the pivot slide into collaborative filtering. Everything "
                  "from here forward in the course is really just different strategies for answering this one "
                  "question.")

# ===========================================================================
# SLIDE 13 — THE SIMPLE IDEA
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC7, "The Simple Idea", "Slide 13")
add_text(s, LEFT_X, Inches(1.75), FULL_W, Inches(0.6), "“People who behave similarly may like similar things.”",
          size=20, color=INK, bold=True)
add_box_with_text(s, LEFT_X, Inches(2.85), Inches(5.6), Inches(2.0), "John likes:\n\nMovie A\nMovie B",
                   fill=BLUE, size=17, radius=0.08)
add_box_with_text(s, Inches(6.85), Inches(2.85), Inches(5.6), Inches(2.0), "Mary likes:\n\nMovie A\nMovie B\nMovie C",
                   fill=TEAL, size=17, radius=0.08)
add_pill(s, Inches(3.0), Inches(5.15), Inches(7.3), Inches(0.65), "Possible recommendation: Movie C → John",
          fill=BLUE_SOFT, text_color=BLUE_DARK, size=16)
add_text(s, LEFT_X, Inches(6.05), FULL_W, Inches(0.5), "This idea is called COLLABORATIVE FILTERING.",
          size=16, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
add_notes(s, "Only reveal the term 'Collaborative Filtering' at the very end of this slide, after the intuition "
             "has landed — this matches the course's teaching philosophy exactly (simple idea before technical "
             "term). This is THE example that recurs throughout the entire course.")

# ===========================================================================
# SLIDE 14 — WHY IS IT CALLED COLLABORATIVE?
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC7, "Why Is It Called Collaborative?", "Slide 14")
add_text(s, LEFT_X, Inches(1.9), FULL_W, Inches(0.8),
          "“The system learns from the combined behaviour\nof many users.”", size=20, color=INK, bold=True)
add_text(s, LEFT_X, Inches(3.2), FULL_W, Inches(0.5),
          "It does not need to understand the movie description first. Instead, it learns from:", size=16, color=SUBTEXT)
add_flow_vertical(s, ["User Behaviour", "+", "Other Users' Behaviour"], Inches(4.0), Inches(4.0), Inches(5.3),
                   box_h=Inches(0.55), gap=Inches(0.18), fill=PURPLE)
add_notes(s, "This directly contrasts collaborative filtering with content-based recommendation (which DOES "
             "need to understand what a movie is 'about' — genre, actors, keywords). We're deliberately not "
             "teaching content-based methods today — mention it exists, briefly, so learners know it's a "
             "different, complementary approach, not that collaborative filtering is the only technique.")

# ===========================================================================
# SLIDE 15 — FIND SIMILAR USERS
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC8, "Find Similar Users", "Slide 15")
add_text(s, LEFT_X, Inches(1.8), FULL_W, Inches(0.8),
          "If two users have similar preferences, we may use one person's behaviour\nto recommend items to the other.",
          size=17, color=INK, bold=True)
add_box_with_text(s, LEFT_X, Inches(2.95), Inches(5.6), Inches(2.1), "John\n\nMovie A → ❤️\nMovie B → ❤️",
                   fill=BLUE, size=17, radius=0.08)
add_box_with_text(s, Inches(6.85), Inches(2.95), Inches(5.6), Inches(2.1),
                   "Mary\n\nMovie A → ❤️\nMovie B → ❤️\nMovie C → ❤️", fill=TEAL, size=17, radius=0.08)
add_pill(s, Inches(4.3), Inches(5.35), Inches(4.7), Inches(0.65), "Recommendation: Movie C → John",
          fill=BLUE_SOFT, text_color=BLUE_DARK, size=15.5)
add_notes(s, "This is user-based collaborative filtering, made concrete — 'people like you also liked...' is "
             "the plain-English summary worth writing on a whiteboard. We'll build this exact mechanic in "
             "Python in Hour 2, using cosine similarity to find users like John.")

# ===========================================================================
# SLIDE 16 — USER-BASED FILTERING WORKFLOW
# ===========================================================================
slide_flow_v(SEC8, "User-Based Filtering Workflow", "Slide 16",
             ["User", "Find Similar Users", "Look at What They Like", "Find Items the User Has Not Seen",
              "Make Recommendations"],
             box_w=Inches(6.0),
             notes="This workflow is exactly what the recommend_movies() function will do in Notebook 02 — "
                   "point that forward-reference out explicitly so learners recognize the code later as a "
                   "direct translation of this diagram, not something new.")

# ===========================================================================
# SLIDE 17 — FIND SIMILAR ITEMS
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC9, "Find Similar Items", "Slide 17")
add_text(s, LEFT_X, Inches(1.8), FULL_W, Inches(1.0),
          "Instead of asking “Which users are similar?” we can ask:\n“Which items are similar based on user behaviour?”",
          size=17.5, color=INK, bold=True)
add_text(s, LEFT_X, Inches(3.3), FULL_W, Inches(0.5), "Users who like Movie A also frequently like Movie B.", size=16.5, color=SUBTEXT)
add_flow_horizontal(s, ["Movie A", "Movie B"], Inches(3.9), Inches(4.2), Inches(5.5), h=Inches(0.9), fill=TEAL)
add_text(s, LEFT_X, Inches(5.5), FULL_W, Inches(0.5), "...may be considered similar.", size=16, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
add_notes(s, "This is the conceptual flip from Slide 15 — instead of comparing PEOPLE, we compare ITEMS based "
             "on which people rated them similarly. This is exactly how our item-based similarity bar chart "
             "(from the notebook) found that 'Steel Vendetta' is most similar to other Action movies — not "
             "because a human tagged them as Action, but purely from co-rating patterns.")

# ===========================================================================
# SLIDE 18 — ITEM-BASED FILTERING WORKFLOW
# ===========================================================================
slide_flow_v(SEC9, "Item-Based Filtering Workflow", "Slide 18",
             ["Item", "Find Similar Items", "Look at User History", "Find Items Matching User History",
              "Make Recommendation"],
             box_w=Inches(6.0),
             notes="Notice the structural symmetry with Slide 16 — same shape, different axis (items instead of "
                   "users). This parallel structure is worth pointing out explicitly; it's why the two "
                   "approaches feel so similar in code, just with the matrix compared along a different "
                   "dimension (rows vs. columns).")

# ===========================================================================
# SLIDE 19 — COMPARISON
# ===========================================================================
slide_table(SEC10, "User-Based vs. Item-Based", "Slide 19",
            ["User-Based", "Item-Based"],
            [["Finds similar users", "Finds similar items"],
             ["Uses user behaviour", "Uses item relationships"],
             ["“People like you...”", "“People who liked this...”"]],
            col_widths=[1, 1], table_h=Inches(2.6),
            notes="This table is the single-slide summary of the whole first half of the course — consider "
                  "leaving it up a little longer while taking questions. Common question: 'which one is "
                  "better?' Honest answer: it depends on the platform — item-based tends to be more stable "
                  "(items don't change taste overnight the way trending user bases can), which is part of why "
                  "Amazon popularized item-based filtering for e-commerce specifically.")

# ===========================================================================
# SLIDE 20 — WHAT DOES SIMILARITY MEAN?
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC11, "What Does Similarity Mean?", "Slide 20")
add_text(s, LEFT_X, Inches(1.9), FULL_W, Inches(0.6),
          "“Similarity measures how closely two users or two items behave.”", size=19, color=INK, bold=True)
add_box_with_text(s, Inches(2.9), Inches(3.1), Inches(7.5), Inches(2.4),
                   "John and Mary\n\nBoth like:\nMovie A, Movie B\n\nSimilarity = High", fill=BLUE, size=18, radius=0.08)
add_notes(s, "Keep this abstract-but-friendly — the concrete mechanics (cosine similarity) come next. The word "
             "'similarity' will map directly onto a NUMBER in the next slide and in the notebooks — a value "
             "between -1 and 1 (in practice, usually 0 to 1 for rating data), where higher means more alike.")

# ===========================================================================
# SLIDE 21 — COSINE SIMILARITY
# ===========================================================================
slide_image_right(SEC11, "Cosine Similarity", "Slide 21",
                   ["Imagine each user's preferences as a DIRECTION.",
                    "If the directions point roughly the same way → HIGH similarity.",
                    "If the directions point very differently → LOW similarity.",
                    "We won't dwell on the formula — the intuition is what matters. (For the curious: "
                    "cosine similarity is the cosine of the angle between two preference vectors.)"],
                   "06_cosine_similarity_intuition.png", img_width=Inches(5.3),
                   notes="Point at the picture: John and Mary's arrows point almost the same direction (small "
                         "angle → high similarity), while David's arrow points somewhere completely different "
                         "(large angle → low similarity). This is literally what sklearn's cosine_similarity() "
                         "computes when we feed it rows of the user-item matrix in Hour 2.")

# ===========================================================================
# SLIDE 22 — THE RECOMMENDATION PIPELINE
# ===========================================================================
slide_flow_v(SEC12, "The Recommendation Pipeline", "Slide 22",
             ["User Data", "User-Item Interactions", "User-Item Matrix", "Calculate Similarity",
              "Find Similar Users or Items", "Generate Scores", "Rank Items", "Recommend Top Items"],
             box_w=Inches(6.0),
             notes="This is the master workflow slide for the whole course — everything maps onto one of these "
                   "eight boxes. Worth returning to a simplified version of this at the very end (Slide 29) so "
                   "learners see the full circle from raw data to a ranked recommendation list.")

# ===========================================================================
# SLIDE 23 — EXAMPLE
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC13, "Making a Recommendation — Example", "Slide 23")
add_bullets(s, LEFT_X, Inches(1.85), FULL_W, Inches(1.8),
            ["John has watched: Movie A, Movie B", "Users similar to John also watched: Movie C, Movie D"],
            size=18)
add_box_with_text(s, Inches(3.4), Inches(3.6), Inches(6.9), Inches(1.6),
                   "1. Movie C\n2. Movie D", fill=BLUE, size=20, radius=0.08)
add_text(s, LEFT_X, Inches(5.5), FULL_W, Inches(0.9),
          "We usually recommend the highest-ranked items that the user has NOT already interacted with.",
          size=16, color=SUBTEXT, italic=True)
add_notes(s, "This 'exclude already-seen items' rule is easy to forget but essential — ask: 'why would it be a "
             "bad idea to recommend John a movie he's already watched and rated?' This exact filtering step "
             "appears explicitly in our recommend_movies() function in Notebook 02.")

# ===========================================================================
# SLIDE 24 — HOW DO WE KNOW OUR RECOMMENDATIONS ARE GOOD?
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC14, "How Do We Know Our Recommendations Are Good?", "Slide 24")
add_box_with_text(s, LEFT_X, Inches(1.9), Inches(5.6), Inches(2.3),
                   "PRECISION\n\n“Out of the items we\nrecommended, how many\nwere relevant?”", fill=BLUE, size=17, radius=0.08)
add_box_with_text(s, Inches(6.85), Inches(1.9), Inches(5.6), Inches(2.3),
                   "RECALL\n\n“Out of the relevant items,\nhow many did we\nsuccessfully recommend?”", fill=TEAL, size=17, radius=0.08)
add_notes(s, "Use a simple analogy: imagine a teacher recommending 5 books to a student. Precision asks 'how "
             "many of those 5 did the student actually enjoy?' Recall asks 'out of ALL the books that student "
             "would have enjoyed, how many did we manage to suggest?' You can have high precision with low "
             "recall (a few perfect picks, but missed many other good ones) or vice versa.")

# ===========================================================================
# SLIDE 25 — TOP K RECOMMENDATIONS
# ===========================================================================
slide_image_right(SEC14, "Top-K Recommendations", "Slide 25",
                   ["In practice, we only show a TOP-K list (e.g. the top 5 recommendations), not every "
                    "possible item.",
                    "Precision@5 = of the 5 items we recommended, how many were relevant?",
                    "Recall@5 = of all the relevant items, how many appeared in our top 5?",
                    "In this example: 3 of 5 recommended items were relevant, so Precision@5 = 0.6."],
                   "09_precision_at_k.png",
                   notes="Keep the math informal, as instructed — the picture does the explaining. Note the "
                         "general trade-off pattern (mention only if there's time): a bigger K tends to raise "
                         "recall (more chances to include relevant items) but often lowers precision (the list "
                         "gets diluted with less-relevant picks).")

# ===========================================================================
# SLIDE 26 — CHALLENGES WITH COLLABORATIVE FILTERING
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC15, "Challenges With Collaborative Filtering", "Slide 26")
add_box_with_text(s, LEFT_X, Inches(1.85), Inches(3.75), Inches(2.3),
                   "COLD START\n\nNew users or items\nhave little or\nno history.", fill=AMBER, size=15.5, radius=0.08)
add_box_with_text(s, Inches(4.78), Inches(1.85), Inches(3.75), Inches(2.3),
                   "SPARSE DATA\n\nMost users interact\nwith only a small\nnumber of items.", fill=RED, size=15.5, radius=0.08)
add_box_with_text(s, Inches(8.55), Inches(1.85), Inches(4.0), Inches(2.3),
                   "SCALABILITY\n\nLarge platforms have\nmillions of users\nand items.", fill=PURPLE, size=15.5, radius=0.08)
add_picture_framed(s, V + "05_full_sparsity_heatmap.png", Inches(4.65), Inches(4.35), w=Inches(4.0))
add_notes(s, "The sparsity heatmap below is from OUR OWN course dataset — 78.7% of the user-movie matrix is "
             "empty, on just 300 users and 60 movies. Emphasize: real platforms are FAR sparser than this "
             "(sometimes 99.9%+ empty), which is exactly why cold start and sparsity are such persistent, "
             "real challenges, not academic curiosities.")

# ===========================================================================
# SLIDE 27 — COMMON MISTAKES
# ===========================================================================
slide_bullets(SEC16, "Common Mistakes", "Slide 27",
              ["Similar users are not always identical.", "Recommendations are predictions, not guarantees.",
               "A recommendation does not mean the user will definitely like the item.",
               "More complex models are not always better.", "User behaviour can change over time."],
              size=18,
              notes="Quick-fire slide — five short, important corrections. If time is short, this is one of the "
                    "safest slides to read briskly, since a fuller misconceptions table appears again near the "
                    "end for instructor reference and the quiz.")

# ===========================================================================
# SLIDE 28 — WHY DO BUSINESSES USE RECOMMENDATION SYSTEMS?
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC17, "Why Do Businesses Use Recommendation Systems?", "Slide 28")
add_bullets(s, COL1_X, Inches(1.85), COL1_W, Inches(3.4),
            ["Better user experience", "Personalization", "Product discovery",
             "Increased engagement", "Increased sales", "Customer retention"], size=17)
add_text(s, COL2_X, Inches(1.85), COL2_W, Inches(0.4), "Used across industries:", size=15, color=MUTED, bold=True)
for i, ind in enumerate(["Streaming", "E-commerce", "Music", "Education", "Banking"]):
    add_pill(s, COL2_X, Inches(2.3 + i * 0.65), Inches(4.2), Inches(0.5), ind, fill=TEAL_SOFT,
              text_color=RGBColor(0x0F, 0x5C, 0x54), size=14)
add_notes(s, "This is the payoff slide — recommendation quality only matters once it changes business outcomes. "
             "Ask: 'can anyone think of a bank or education example?' (Banking: recommending relevant financial "
             "products; Education: recommending the next course a student is likely to complete successfully.)")

# ===========================================================================
# SLIDE 29 — THE COMPLETE MENTAL MODEL
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, SEC18, "The Complete Mental Model", "Slide 29")
add_flow_horizontal(s, ["Users + Items", "Interactions"], LEFT_X, Inches(1.75), FULL_W, h=Inches(0.75), fill=BLUE)
steps2 = ["User-Item Matrix", "Find Similarity", "User-Based or Item-Based", "Predict Interest",
          "Rank Items", "Top Recommendations"]
x = LEFT_X
y = Inches(2.85)
box_w = Inches(5.7)
add_flow_vertical(s, steps2, Inches((13.333-5.7)/2), y, box_w, box_h=Inches(0.46), gap=Inches(0.10), fill=TEAL)
add_notes(s, "This is the single-slide summary of the ENTIRE conceptual hour — consider lingering here for "
             "questions before the coding break. Every slide in Hour 1 maps to one box in this chain.")

# ===========================================================================
# SLIDE 30 — TRANSITION TO CODING
# ===========================================================================
slide_quote(SEC18, "Transition to Coding", "Slide 30",
            "“Now that we understand how collaborative filtering works,\nhow do we build a simple recommendation engine in Python?”",
            subtext="Let's open the notebooks.",
            notes="Short, energetic transition slide. Take a 2-3 minute break here if the room needs it before "
                  "diving into Hour 2's hands-on notebooks.")

# ===========================================================================
# SLIDE 31 — IMPORTANT DISTINCTIONS
# ===========================================================================
slide_table("Instructor Reference", "Important Distinctions", "Slide 31",
            ["Myth", "Reality"],
            [["A recommendation system knows exactly what users want.", "It PREDICTS what users may be "
              "interested in — it does not know for certain."],
             ["Similar users are exactly the same.", "Similarity means users have SOME patterns in common, "
              "not identical taste."],
             ["A recommendation guarantees the user will like something.", "Recommendations are predictions, "
              "with uncertainty."],
             ["Missing ratings mean the user dislikes the item.", "Missing information usually just means the "
              "user has not interacted with that item yet."],
             ["More data automatically solves the cold-start problem.", "A brand-new user can still have zero "
              "history, no matter how much data exists for OTHER users."]],
            col_widths=[1, 1.5], table_h=Inches(4.8),
            notes="Consider running this as a rapid true/false check with the class before the quiz — these "
                  "five distinctions are the ones beginners misstate most often in follow-up conversations.")

# ===========================================================================
# SLIDE 32 — QUICK QUIZ
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, "Knowledge Check", "Quick Quiz", "Slide 32")
quiz = ["1. What is a recommender system?", "2. What is the difference between explicit and implicit feedback?",
        "3. What is a user-item matrix?", "4. What is collaborative filtering?",
        "5. What is user-based collaborative filtering?", "6. What is item-based collaborative filtering?",
        "7. What is the cold-start problem?", "8. Does a recommendation guarantee that a user will like an item?"]
add_bullets(s, COL1_X, CONTENT_TOP, COL1_W, Inches(5.3), quiz[:4], size=16, space_after=18, marker="")
add_bullets(s, COL2_X, CONTENT_TOP, COL2_W, Inches(5.3), quiz[4:], size=16, space_after=18, marker="")
add_notes(s, "ANSWER KEY (for instructor use):\n"
             "1. A system that helps users discover items they may be interested in, based on patterns in "
             "behaviour rather than explicit rules.\n"
             "2. Explicit feedback is when a user directly tells us their preference (a star rating, a "
             "like/dislike). Implicit feedback is inferred from behaviour (clicks, purchases, watch history) "
             "without the user directly stating an opinion.\n"
             "3. A table where rows are users, columns are items, and values are ratings or interactions — "
             "missing entries represent unknown preferences.\n"
             "4. Using patterns in the combined behaviour of many users to make recommendations, without "
             "needing to understand item descriptions or content.\n"
             "5. Finding users with similar preferences to a target user, and recommending items those similar "
             "users liked that the target user hasn't seen yet.\n"
             "6. Finding items that tend to be liked by the same users (similar rating patterns), and "
             "recommending items similar to ones a user already likes.\n"
             "7. The problem of making good recommendations for new users or new items that have little or no "
             "interaction history yet.\n"
             "8. No. A recommendation is a PREDICTION of possible interest, not a guarantee — the user may "
             "still dislike it.")

# ===========================================================================
# SLIDE 33 — REFERENCES / SOURCES
# ===========================================================================
s = blank_slide(prs)
add_kicker_title(s, "Appendix", "References / Sources", "Slide 33")
add_text(s, LEFT_X, CONTENT_TOP, FULL_W, Inches(0.32), "Textbook", size=14, color=MUTED, bold=True)
add_bullets(s, LEFT_X, Inches(1.92), FULL_W, Inches(1.2),
            ["Aggarwal, C. C. (2015). Data Mining: The Textbook. Springer. — Section 18.5, “Recommender "
             "Systems” (utility matrices, neighborhood-based user-based and item-based collaborative filtering)."],
            size=14, space_after=6)
add_text(s, LEFT_X, Inches(3.1), FULL_W, Inches(0.32), "Official Documentation (verified for current syntax)",
          size=14, color=MUTED, bold=True)
add_bullets(s, LEFT_X, Inches(3.47), FULL_W, Inches(1.7),
            ["scikit-learn: sklearn.metrics.pairwise.cosine_similarity "
             "(scikit-learn.org/stable/modules/generated/ ...)",
             "pandas API Reference: DataFrame.pivot_table() (pandas.pydata.org/docs/reference/ ...)",
             "Every code snippet in this deck and the notebooks was checked against these pages during "
             "development."],
            size=13, space_after=6)
add_text(s, LEFT_X, Inches(5.35), FULL_W, Inches(0.32), "Evaluation Concepts", size=14, color=MUTED, bold=True)
add_bullets(s, LEFT_X, Inches(5.72), FULL_W, Inches(0.9),
            ["Precision@K and Recall@K definitions cross-checked against standard recommender-systems "
             "evaluation references (not covered explicitly in the primary textbook)."],
            size=13, space_after=6)
add_text(s, LEFT_X, Inches(6.7), FULL_W, Inches(0.32), "Data", size=14, color=MUTED, bold=True)
add_bullets(s, LEFT_X, Inches(7.02), FULL_W, Inches(0.4),
            ["All ratings data used in this course is SYNTHETIC (data/generate_data.py). Movie titles are "
             "entirely invented."], size=11.5, space_after=4)
add_notes(s, "All scikit-learn and pandas syntax in this deck (cosine_similarity, pivot_table) was checked "
             "against the official documentation during course development to ensure it reflects current, "
             "non-deprecated APIs.")

# ===========================================================================
# Final pass: footers
# ===========================================================================
total = len(prs.slides)
for i, sl in enumerate(prs.slides, start=1):
    add_footer(sl, i, total, tag="Collaborative Filtering Engines")

prs.save("Collaborative_Filtering_Engines_Beginner_Course.pptx")
print(f"Saved deck with {total} slides.")
